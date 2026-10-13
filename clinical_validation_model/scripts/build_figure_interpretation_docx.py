from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x20, 0x2A, 0x35)
MUTED = RGBColor(0x5A, 0x65, 0x70)
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
EXPECTED_FIGURE_IDS = tuple(f"{index:02d}" for index in range(1, 13))
FIELD_LABELS = (
    "Purpose",
    "Comparison",
    "Observed result",
    "Preferred direction",
    "Interpretation caution",
    "Improvement implication",
    "Current limitation",
)


def _set_run_font(
    run,
    *,
    size: float,
    color: RGBColor = INK,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def _configure_style(style, *, size: float, color: RGBColor, before: float, after: float) -> None:
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.keep_with_next = True


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    _configure_style(doc.styles["Heading 1"], size=16, color=BLUE, before=18, after=10)
    _configure_style(doc.styles["Heading 2"], size=13, color=BLUE, before=14, after=7)
    _configure_style(doc.styles["Heading 3"], size=12, color=DARK_BLUE, before=10, after=5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    _set_run_font(
        header.add_run("SiBC Clinical Validation Model | Release Technical Report"),
        size=8.5,
        color=MUTED,
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    _set_run_font(footer.add_run("Page "), size=8.5, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def _load_manifest(run_root: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    manifest = json.loads(
        (run_root / "artifacts" / "reproducibility_manifest.json").read_text(
            encoding="utf-8"
        )
    )
    figures = json.loads(
        (run_root / "figures" / "figure_manifest.json").read_text(encoding="utf-8")
    )
    if not isinstance(manifest, dict) or not isinstance(figures, list):
        raise ValueError("Invalid release report manifest contract")
    return manifest, figures


def _parse_interpretation(path: Path) -> dict[str, dict[str, str]]:
    text = path.read_text(encoding="utf-8")
    matches = list(
        re.finditer(
            r"^## Figure (?P<id>\d{2})\. (?P<title>.+)$",
            text,
            flags=re.MULTILINE,
        )
    )
    sections: dict[str, dict[str, str]] = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        values = {"Title": match.group("title").strip()}
        for line in text[match.end() : end].splitlines():
            if not line.startswith("- ") or ": " not in line:
                continue
            label, value = line[2:].split(": ", 1)
            values[label.strip()] = value.strip()
        sections[match.group("id")] = values
    if tuple(sorted(sections)) != EXPECTED_FIGURE_IDS:
        raise ValueError("Figure interpretation guide must contain Figures 01-12")
    for figure_id, values in sections.items():
        missing = [label for label in FIELD_LABELS if not values.get(label)]
        if missing:
            raise ValueError(f"Figure {figure_id} is missing interpretation fields: {missing}")
    return sections


def _png_by_figure(figures: list[dict[str, Any]], run_root: Path) -> dict[str, Path]:
    mapping: dict[str, Path] = {}
    for record in figures:
        filename = str(record.get("file_name", ""))
        figure_id = str(record.get("figure_id", ""))
        if filename.endswith(".png") and figure_id in EXPECTED_FIGURE_IDS:
            path = (run_root / "figures" / filename).resolve()
            if path.parent != (run_root / "figures").resolve() or not path.is_file():
                raise ValueError("Figure manifest contains an unmanaged PNG path")
            mapping[figure_id] = path
    if tuple(sorted(mapping)) != EXPECTED_FIGURE_IDS:
        raise ValueError("Figure manifest must reference exactly 12 governed PNG files")
    return mapping


def _add_cover(doc: Document, manifest: dict[str, Any]) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    _set_run_font(kicker.add_run("RELEASE TECHNICAL EVALUATION"), size=10, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    _set_run_font(title.add_run("Clinical Validation Model"), size=28, color=DARK_BLUE, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    _set_run_font(subtitle.add_run("Figure Interpretation Guide"), size=16, color=BLUE)

    boundary = doc.add_paragraph()
    boundary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    boundary.paragraph_format.space_after = Pt(34)
    _set_run_font(
        boundary.add_run(
            "Synthetic scenario-based technical evaluation only. "
            "Clinical performance remains BLOCKED_BY_REFERENCE_STANDARD."
        ),
        size=11,
        color=INK,
        italic=True,
    )
    metadata = (
        ("Release run", str(manifest.get("run_id", "NOT_AVAILABLE"))),
        ("Model version", str(manifest.get("model_version", "NOT_AVAILABLE"))),
        ("Release snapshot", str(manifest.get("data_snapshot_id", "NOT_AVAILABLE"))),
        ("Execution mode", str(manifest.get("execution_mode", "release_artifact"))),
        ("Clinical validation", "BLOCKED_BY_REFERENCE_STANDARD"),
    )
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        _set_run_font(paragraph.add_run(f"{label}: "), size=10.5, color=MUTED, bold=True)
        _set_run_font(paragraph.add_run(value), size=10.5, color=INK)


def _add_figure_page(
    doc: Document,
    figure_id: str,
    values: dict[str, str],
    image_path: Path,
) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(7)
    heading.paragraph_format.keep_with_next = True
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run_font(
        heading.add_run(f"Figure {figure_id}. {values['Title']}"),
        size=14,
        color=BLUE,
        bold=True,
    )
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_after = Pt(2)
    picture_paragraph.add_run().add_picture(str(image_path), width=Inches(5.75))
    drawing = picture_paragraph._p.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", f"Figure {figure_id}: {values['Title']}")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)
    _set_run_font(
        caption.add_run("Synthetic scenario-based technical evaluation"),
        size=8,
        color=MUTED,
        italic=True,
    )
    for label in FIELD_LABELS:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.paragraph_format.keep_together = True
        _set_run_font(paragraph.add_run(f"{label}: "), size=9, color=DARK_BLUE, bold=True)
        _set_run_font(paragraph.add_run(values[label]), size=9, color=INK)


def _audit_document(doc: Document) -> None:
    section = doc.sections[0]
    expected = {
        "page_width": Inches(PAGE_WIDTH_IN),
        "page_height": Inches(PAGE_HEIGHT_IN),
        "top_margin": Inches(MARGIN_IN),
        "right_margin": Inches(MARGIN_IN),
        "bottom_margin": Inches(MARGIN_IN),
        "left_margin": Inches(MARGIN_IN),
    }
    for attribute, value in expected.items():
        if getattr(section, attribute) != value:
            raise ValueError(f"DOCX preset audit failed: {attribute}")
    style_tokens = {
        "Normal": (11, 0, 6),
        "Heading 1": (16, 18, 10),
        "Heading 2": (13, 14, 7),
        "Heading 3": (12, 10, 5),
    }
    for style_name, (size, before, after) in style_tokens.items():
        style = doc.styles[style_name]
        if style.font.name != "Calibri" or style.font.size != Pt(size):
            raise ValueError(f"DOCX preset audit failed: {style_name} typography")
        if (
            style.paragraph_format.space_before != Pt(before)
            or style.paragraph_format.space_after != Pt(after)
            or style.paragraph_format.line_spacing != 1.25
        ):
            raise ValueError(f"DOCX preset audit failed: {style_name} spacing")
    if len(doc.inline_shapes) != 12:
        raise ValueError("DOCX preset audit failed: expected 12 inline figures")
    for table in doc.tables:
        raise ValueError(f"DOCX preset audit failed: unexpected table {table}")


def build_figure_interpretation_docx(run_root: Path) -> Path:
    run_root = run_root.resolve()
    if run_root.parent.name != "release":
        raise ValueError("DOCX output is restricted to a release run directory")
    manifest, figures = _load_manifest(run_root)
    sections = _parse_interpretation(
        run_root / "reports" / "figure_interpretation_guide.md"
    )
    images = _png_by_figure(figures, run_root)
    doc = Document()
    _configure_document(doc)
    doc.core_properties.title = "SiBC Clinical Validation Model - Figure Interpretation Guide"
    doc.core_properties.subject = "Synthetic scenario-based release technical evaluation"
    doc.core_properties.author = "SiBC Clinical Validation Model"
    doc.core_properties.last_modified_by = "SiBC Clinical Validation Model"
    doc.core_properties.keywords = "SiBC, clinical validation, synthetic technical evaluation"
    doc.core_properties.comments = "No clinical performance claim."
    created = datetime.now(timezone.utc).replace(microsecond=0)
    doc.core_properties.created = created
    doc.core_properties.modified = created
    _add_cover(doc, manifest)
    for figure_id in EXPECTED_FIGURE_IDS:
        _add_figure_page(doc, figure_id, sections[figure_id], images[figure_id])
    _audit_document(doc)
    destination = run_root / "reports" / "figure_interpretation_guide.docx"
    doc.save(destination)
    return destination


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build the governed release figure interpretation DOCX."
    )
    parser.add_argument("--run-root", required=True, type=Path)
    args = parser.parse_args(argv)
    destination = build_figure_interpretation_docx(args.run_root)
    print(destination)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
